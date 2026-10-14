"""Stage 1: parse Telegram Desktop HTML exports into message records.

XPath notes (adapted from the task brief after inspecting real fixtures in
analysis_tools/tests/fixtures/):
- The per-message timestamp div is class="pull_right date details" — two
  tokens, NOT "date_details" as the brief guessed — so we match divs having
  both "date" and "details". The forwarded-original date inside from_name is a
  <span>, never a div, so div-only matching cannot pick it up by accident.
- Forwarded messages carry a <div class="pull_left forwarded userpic_wrap">
  BEFORE the real origin container <div class="forwarded body">. Matching bare
  contains(@class,"forwarded") hits the userpic wrapper first, so we require
  both "forwarded" and "body".
- The origin name sits in from_name INSIDE the forwarded body, with a nested
  <span class="date details">original date</span> whose text must be excluded.
- sender is the OUTER from_name (who posted in this chat); forwarded_from is
  the INNER from_name (original author). The brief's draft made sender equal
  the forwarded origin, which loses the poster.
- No span[class~="user"] fallback: in these exports the only span classes
  starting with "user" are reaction "userpics", which would pollute the origin
  with avatar initials.
- reply_to hrefs look like "#go_to_message802"; we strip the "go_to_" marker so
  the value matches the message_id format ("message802") for joins.

Fix round 1:
- Message selection now uses starts-with(normalize-space(@class),"message"):
  a bare contains(@class,"message") also matched NESTED divs such as
  <div class="media clearfix pull_left media_voice_message"> and
  <div class="text rich_message">, producing phantom nosid- records (proven:
  1112 records vs 1068 real message divs in private_group/messages2.html).
  All three real message variants ("message default clearfix[ joined]",
  "message service") start with "message"; no nested class does.
- Media kinds map (verified across all 19 export HTML files): only
  photo_wrap (2653x), media_poll (7x) and media_voice_message (314x) occur;
  the brief's other guesses (video_file_wrap, voice_player, animated_wrap,
  sticker_wrap, document_wrap, media_audio_player) occur 0x and are pruned.
  Voice blocks use media_voice_message -> "voice".
"""
import re
from datetime import datetime
from pathlib import Path
from lxml import html as lhtml

DATE_RE = re.compile(r"(\d{2})\.(\d{2})\.(\d{4}) (\d{2}):(\d{2}):(\d{2})(?:\s*(UTC[+-][\d:]+))?")
# class token -> media kind; only classes verified present in the real exports
MEDIA_CLASSES = {"photo_wrap": "photo", "media_voice_message": "voice", "media_poll": "poll"}

def parse_timestamp(title):
    m = DATE_RE.search(title or "")
    if not m:
        return None, None
    d, mo, y, h, mi, s, tz = m.groups()
    return datetime(int(y), int(mo), int(d), int(h), int(mi), int(s)), tz

def iter_message_files(chat_dir: Path):
    return sorted(chat_dir.glob("messages*.html"))

def _media_kinds(node):
    out = []
    for cls, kind in MEDIA_CLASSES.items():
        if node.xpath(f'.//*[contains(@class,"{cls}")]'):
            out.append(kind)
    return out

def _clean(texts):
    return re.sub(r"\s+", " ", " ".join(t.strip() for t in texts if t and t.strip())).strip()

# text nodes of a from_name, excluding the nested original-date span
_FROM_NAME_TEXT = ('.//div[contains(@class,"from_name")]'
                   '//text()[not(ancestor::span[contains(@class,"date")])'
                   ' and not(ancestor::span[contains(@class,"details")])]')

def parse_file(path: Path, chat: str):
    doc = lhtml.parse(str(path))
    records = []
    for msg in doc.xpath('//div[starts-with(normalize-space(@class),"message")]'):
        mid = msg.get("id", "")
        titles = msg.xpath('.//div[contains(@class,"date") and contains(@class,"details")]/@title')
        ts, tz = parse_timestamp(titles[0]) if titles else (None, None)
        is_service = "service" in (msg.get("class") or "")
        rec = {"chat": chat, "message_id": mid or f"nosid-{len(records)}",
               "timestamp": ts.isoformat() if ts else None, "tz": tz,
               "sender": None, "forwarded_from": None, "text_plain": "",
               "media_kinds": _media_kinds(msg), "is_reply": False,
               "reply_to": None, "is_service": is_service}
        if is_service:
            rec["text_plain"] = _clean(msg.xpath("./text() | .//text()"))
            records.append(rec)
            continue
        fwd = msg.xpath('.//div[contains(@class,"forwarded") and contains(@class,"body")]')
        if fwd:
            origin = fwd[0].xpath(_FROM_NAME_TEXT)
            rec["forwarded_from"] = _clean(origin) or None
        names = msg.xpath('.//div[contains(@class,"from_name")]'
                          '[not(ancestor::div[contains(@class,"forwarded")])]'
                          '//text()[not(ancestor::span[contains(@class,"date")])'
                          ' and not(ancestor::span[contains(@class,"details")])]')
        rec["sender"] = _clean(names) or None
        reply = msg.xpath('.//div[contains(@class,"reply_to")]//a/@href')
        if reply:
            rec["is_reply"] = True
            rec["reply_to"] = re.sub(r"^go_to_", "", (reply[0] or "").lstrip("#")) or None
        body_text = msg.xpath('.//div[contains(@class,"text")]')
        rec["text_plain"] = _clean([" ".join(b.itertext()) for b in body_text])
        records.append(rec)
    return records

import jdatetime
from .common import ROOT, SOURCES, write_json, WINDOW_START, WINDOW_END

# Persian weekday names only render when the date's locale is FA_LOCALE;
# without it jd.strftime("%A") returns the English name on non-Persian hosts.
_FA_LOCALE = getattr(jdatetime, "FA_LOCALE", "Persian_Iran")

def enrich(rec):
    if rec.get("timestamp"):
        dt = datetime.fromisoformat(rec["timestamp"])
        jd = jdatetime.date.fromgregorian(date=dt.date(), locale=_FA_LOCALE)
        rec["date_greg"] = dt.date().isoformat()
        rec["jalali_month"] = f"{jd.year}/{jd.month:02d}"
        rec["jalali_weekday"] = jd.strftime("%A")
        rec["hour"] = dt.hour
        rec["weekday_en"] = dt.strftime("%A")
        rec["in_window"] = WINDOW_START <= rec["date_greg"] <= WINDOW_END
    else:
        for k in ("date_greg", "jalali_month", "jalali_weekday", "hour", "weekday_en"):
            rec[k] = None
        rec["in_window"] = False
    return rec

def parse_docx(path):
    import docx
    d = docx.Document(str(path))
    paras = [{"style": p.style.name if p.style else "Normal", "text": p.text}
             for p in d.paragraphs if p.text.strip()]
    tables = [[cell.text.strip() for cell in row.cells] for t in d.tables for row in t.rows]
    return {"paragraphs": paras, "tables": tables}

def run_all():
    import json
    out = []
    for chat, src in SOURCES.items():
        for f in iter_message_files(src):
            for rec in parse_file(f, chat):
                out.append(enrich(rec))
    with open(ROOT / "analysis/messages.jsonl", "w", encoding="utf-8") as fh:
        for rec in out:
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
    write_json(ROOT / "analysis/docx_content.json",
               parse_docx(ROOT / "گزارش_عملکرد_رسانه‌ای_انجمن_علمی_فیزیک.docx"))
    print(f"parsed {len(out)} messages -> analysis/messages.jsonl")

if __name__ == "__main__":
    run_all()
