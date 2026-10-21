# -*- coding: utf-8 -*-
"""Extract text from the governance source documents (Persian PDFs/DOCX).

The sources live in sources/governance/ (git-ignored: the repo is public and
some sources carry personal details). Output: sources/governance/extracts/<slug>.txt

Usage: python -m analysis_tools.extract_governance_sources
"""
import sys
import pathlib

from .common import MEDIA_DOCX, ROOT

sys.stdout.reconfigure(encoding="utf-8")

SRC = ROOT / "sources" / "governance"
OUT = SRC / "extracts"
OUT.mkdir(parents=True, exist_ok=True)

SOURCES = [
    # (path, ascii slug)
    (SRC / "اساس‌نامه انجمن علمی فیزیک.pdf", "asasname_physics_current"),
    (SRC / "مجمع انجمن علمی دانشکده فیزیک.pdf", "majma_charter"),
    (SRC / "Asasname.pdf", "asasname_resana_reference"),
    (SRC / "سند_چارچوب_عملیاتی_انجمن_علمی_دانشکده_فیزیک.pdf", "operational_framework"),
    (SRC / "راهنمای_حقوقی_و_تکالیف_انجمن_علمی_فیزیک_شریف.docx", "legal_guide"),
    (SRC / "دستورجلسه 10مهر.pdf", "agenda_10mehr"),
    (SRC / "گزارش 15 تیر انجمن.pdf", "report_15tir"),
    (SRC / "تیم چای.pdf", "tea_team"),
    (SRC / "چارچوب چای.pdf", "tea_framework"),
    (SRC / "گزارش برنامه.docx", "program_report"),
    (SRC / "Kaveh_Sharif_Physics_Proposal_Revised.docx", "kaveh_proposal"),
    (MEDIA_DOCX, "media_performance_report"),
]


def pdf_text(path: pathlib.Path) -> tuple[str, int]:
    import fitz
    doc = fitz.open(path)
    pages = []
    for i, page in enumerate(doc, 1):
        pages.append(f"===== PAGE {i} =====\n{page.get_text('text')}")
    n = len(doc)
    doc.close()
    return "\n".join(pages), n


def docx_text(path: pathlib.Path) -> tuple[str, int]:
    from docx import Document
    doc = Document(str(path))
    parts = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            style = para.style.name if para.style else ""
            prefix = "## " if style.lower().startswith("heading") else ""
            parts.append(prefix + t)
    for table in doc.tables:
        parts.append("----- TABLE -----")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts), 0


for path, slug in SOURCES:
    if not path.exists():
        print(f"MISSING: {path}")
        continue
    try:
        if path.suffix.lower() == ".pdf":
            text, n = pdf_text(path)
        else:
            text, n = docx_text(path)
        out = OUT / f"{slug}.txt"
        out.write_text(text, encoding="utf-8")
        print(f"OK  {slug}: {n} pages, {len(text)} chars -> {out.name}")
    except Exception as e:
        print(f"ERR {slug}: {e}")
