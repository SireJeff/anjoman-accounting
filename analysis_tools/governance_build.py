# -*- coding: utf-8 -*-
"""Render governance markdown masters to RTL DOCX (+ PDF via report_common).

OOXML notes (Persian is a complex script):
  - A bidi paragraph (<w:bidi/> in pPr) is right-aligned/start-aligned by
    default, so paragraphs here do NOT also set jc; setting both is
    redundant and the brief's original draft appended bidi after jc,
    which violates CT_PPr's defined child order.
  - Word distinguishes "Western" run properties (w:b, w:sz, w:rFonts/@ascii)
    from "complex-script" run properties (w:bCs, w:szCs, w:rFonts/@cs). A
    run flagged rtl (w:rtl, exposed by python-docx as Font.rtl) is shaped
    with the complex-script properties, so bold/size must be mirrored onto
    w:bCs/w:szCs or headings render plain in Word.
"""
from __future__ import annotations

import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt

from . import governance_lint as gl
from .report_common import md_to_pdf_rtl

FONT = "Tahoma"


def _rtl_paragraph(p):
    """Mark a paragraph bidi (RTL). Do not also set jc/alignment: a bidi
    paragraph's default alignment is already right/start, and CT_PPr
    requires its children in schema order (bidi after jc would be invalid)."""
    pPr = p._p.get_or_add_pPr()
    pPr.append(OxmlElement("w:bidi"))
    return p


def _rtl_run(r, bold=False, size=None):
    r.font.name = FONT
    r.font.rtl = True
    if bold:
        r.bold = True
        r.font.cs_bold = True
    if size:
        r.font.size = Pt(size)
    rPr = r._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:cs"), FONT)
    if size:
        # python-docx has no Font property for szCs; w:sz/w:szCs share the
        # same half-point value and CT_RPr requires szCs immediately after
        # sz, so insert it by hand right after the sz element python-docx
        # just wrote via r.font.size.
        sz = rPr.find(qn("w:sz"))
        szCs = OxmlElement("w:szCs")
        if sz is not None:
            szCs.set(qn("w:val"), sz.get(qn("w:val")))
            sz.addnext(szCs)
        else:
            rPr.append(szCs)
    return r


_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _add_rich(p, text, bold=False, size=None):
    pos = 0
    for m in _BOLD.finditer(text):
        if m.start() > pos:
            _rtl_run(p.add_run(text[pos:m.start()]), bold=bold, size=size)
        _rtl_run(p.add_run(m.group(1)), bold=True, size=size)
        pos = m.end()
    if pos < len(text):
        _rtl_run(p.add_run(text[pos:]), bold=bold, size=size)


def md_to_docx_rtl(md_text: str, title: str) -> Document:
    doc = Document()
    doc.core_properties.title = title
    for line in md_text.splitlines():
        s = line.strip()
        if not s:
            continue
        if s.startswith("# ") and not s.startswith("## "):
            p = _rtl_paragraph(doc.add_paragraph())
            _add_rich(p, s[2:], bold=True, size=16)
        elif s.startswith("## "):
            p = _rtl_paragraph(doc.add_paragraph())
            _add_rich(p, s[3:], bold=True, size=13)
        elif s.startswith("### "):
            p = _rtl_paragraph(doc.add_paragraph())
            _add_rich(p, s[4:], bold=True, size=12)
        elif s.startswith("- "):
            p = _rtl_paragraph(doc.add_paragraph())
            _rtl_run(p.add_run("•  "))
            _add_rich(p, s[2:])
        else:
            p = _rtl_paragraph(doc.add_paragraph())
            _add_rich(p, s)
    return doc


def _rezip_deterministic(path: Path) -> None:
    """Rewrite the docx zip with fixed entry timestamps -> byte-stable output."""
    fixed = (1980, 1, 1, 0, 0, 0)
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td) / "rezip.docx"
        with zipfile.ZipFile(path) as zin, \
             zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                info.date_time = fixed
                zout.writestr(info, data)
        shutil.move(str(tmp), str(path))


def build_all(out_dir: Path | None = None) -> list[Path]:
    out = Path(out_dir) if out_dir else gl.GOV / "build"
    out.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for key, md in gl.MASTER_FILES.items():
        if not md.exists():
            print(f"skip {key}: {md} missing", file=sys.stderr)
            continue
        text = md.read_text(encoding="utf-8")
        title = text.splitlines()[0].lstrip("# ").strip()
        stem = md.stem
        docx_path = out / f"{stem}.docx"
        pdf_path = out / f"{stem}.pdf"
        d = md_to_docx_rtl(text, title=title)
        d.save(str(docx_path))
        _rezip_deterministic(docx_path)
        md_to_pdf_rtl(md, str(pdf_path), title=title,
                      footer="انجمن علمی دانشکده فیزیک دانشگاه صنعتی شریف")
        written += [docx_path, pdf_path]
    return written


if __name__ == "__main__":
    for f in build_all():
        print(f)
